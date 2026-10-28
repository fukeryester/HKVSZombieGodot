"""Read editable Excel inputs; authoring is handled by Artifact Tool."""
import json,sys,os
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
result={}
for file in sorted(Path(os.environ.get('CONFIG_DIR',str(root/'config'))).glob('*.xlsx')):
    if file.name.startswith('~$'):continue
    wb=load_workbook(file,read_only=True,data_only=False)
    for row in wb['参数'].iter_rows(min_row=2,values_only=True):
        key,value=row[0],row[2]
        if key is None:continue
        if key in result:raise ValueError('Duplicate configuration key: '+key)
        if isinstance(value,bool) or not isinstance(value,(int,float)):raise ValueError('Numeric value required (no formulas): '+key)
        result[key]=value
print(json.dumps(result,ensure_ascii=True,allow_nan=False))
